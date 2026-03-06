import copy
import logging
import os
import re
from io import BytesIO
from typing import List, Optional, Union

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Pt, Emu, Inches

from ppt_gen.models import DayData, ProposalData, Quote

logger = logging.getLogger(__name__)

TEMPLATE = os.path.join(os.path.dirname(__file__), "SSA PP template.pptx")

ORANGE = RGBColor(0xAC, 0x4E, 0x27)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
DARK   = RGBColor(0x2B, 0x2B, 0x2B)
SAGE   = RGBColor(0x86, 0x96, 0x7F)

FT = "Wulkan Display Medium"
FB = "Inter"

FS_SLIDE_TITLE = 62
FS_SIDEBAR     = 38
FS_HEAD        = 34
FS_BODY        = 28
FS_SMALL       = 22
FS_SAY         = 20
FS_QUOTE       = 20
FS_AMOUNT      = 90
FS_TITLE_MAIN  = 78
FS_TITLE_NAMES = 50

SIDEBAR_W = Emu(2_979_420)
SLIDE_W   = Emu(18_288_000)
SLIDE_H   = Emu(10_287_000)
PHOTO_Y   = Emu(5_781_293)
PAD       = Inches(0.28)

CX  = SIDEBAR_W + PAD
CW  = SLIDE_W - CX - PAD
CY  = Inches(0.45)
CH  = PHOTO_Y - CY - PAD

COL_W  = (CW - PAD) // 2
COL2_X = CX + COL_W + PAD

SL_X = Inches(0.17)
SL_Y = Inches(2.30)
SL_W = SIDEBAR_W - Inches(0.20)
SL_H = Inches(3.50)

def _xml_run(para_el, text: str, font: str, size_pt: int,
             color: RGBColor, bold: bool = False, italic: bool = False) -> None:
    r = etree.SubElement(para_el, qn("a:r"))
    rPr = etree.SubElement(r, qn("a:rPr"))
    rPr.set("lang", "en-US")
    rPr.set("sz", str(size_pt * 100))
    rPr.set("dirty", "0")
    if bold:
        rPr.set("b", "1")
    if italic:
        rPr.set("i", "1")
    fill = etree.SubElement(rPr, qn("a:solidFill"))
    clr  = etree.SubElement(fill, qn("a:srgbClr"))
    clr.set("val", f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
    lat = etree.SubElement(rPr, qn("a:latin"))
    lat.set("typeface", font)
    t = etree.SubElement(r, qn("a:t"))
    t.text = text

def _add_tb(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    return tf

def _tb1(slide, x, y, w, h,
         text: str, font: str, size: int, color: RGBColor,
         bold: bool = False, italic: bool = False,
         align: PP_ALIGN = PP_ALIGN.LEFT) -> None:
    tf = _add_tb(slide, x, y, w, h)
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color

def _run(para, text: str, font: str, size: int, color: RGBColor,
         bold: bool = False, italic: bool = False,
         align: PP_ALIGN = PP_ALIGN.LEFT) -> None:
    para.alignment = align
    r = para.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color


def _remap_rels(src_part, dst_part, element) -> None:
    RREL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    for el in element.iter():
        for attr in (f"{{{RREL}}}embed", f"{{{RREL}}}link", f"{{{RREL}}}id"):
            rId = el.get(attr)
            if rId and rId in src_part.rels:
                rel = src_part.rels[rId]
                new_rId = dst_part.relate_to(rel.target_part, rel.reltype)
                el.set(attr, new_rId)

def _dup_slide(prs, src_idx: int):
    src = prs.slides[src_idx]
    new = prs.slides.add_slide(src.slide_layout)

    spTree = new.shapes._spTree
    for child in list(spTree)[2:]:
        spTree.remove(child)
    for child in list(src.shapes._spTree)[2:]:
        el = copy.deepcopy(child)
        _remap_rels(src.part, new.part, el)
        spTree.append(el)

    src_cSld = src._element.find(qn("p:cSld"))
    dst_cSld = new._element.find(qn("p:cSld"))
    src_bg = src_cSld.find(qn("p:bg"))
    if src_bg is not None:
        old_bg = dst_cSld.find(qn("p:bg"))
        if old_bg is not None:
            dst_cSld.remove(old_bg)
        new_bg = copy.deepcopy(src_bg)
        _remap_rels(src.part, new.part, new_bg)
        dst_cSld.insert(0, new_bg)

    return new


def _s0_title(slide, client_names: str, quotes: Union[Optional[Quote], List[Quote]],
              retreat_title: str = "Transformational\nSoul Coach Training") -> None:
    for shape in slide.shapes:
        if shape.name == "TextBox 7" and shape.has_text_frame:
            shape.left   = Inches(1.75)
            shape.width  = Inches(16.5)
            shape.height = Inches(4.5)

            txBody = shape.text_frame._txBody
            for p in list(txBody.findall(qn("a:p"))):
                txBody.remove(p)

            for line in retreat_title.split("\n"):
                p1 = etree.SubElement(txBody, qn("a:p"))
                pPr1 = etree.SubElement(p1, qn("a:pPr"))
                pPr1.set("algn", "ctr")
                _xml_run(p1, line, FT, 52, ORANGE)

            etree.SubElement(txBody, qn("a:p"))

            p2 = etree.SubElement(txBody, qn("a:p"))
            pPr2 = etree.SubElement(p2, qn("a:pPr"))
            pPr2.set("algn", "ctr")
            _xml_run(p2, client_names, FB, 54, ORANGE)
            break

    quote_list = []
    if quotes:
        if isinstance(quotes, list):
            quote_list = quotes[:2]
        else:
            quote_list = [quotes]
    
    if quote_list:
        if len(quote_list) >= 2:
            quote1 = quote_list[0]
            quote2 = quote_list[1]
            text1 = f'\u201c{quote1.text}\u201d'
            if quote1.author:
                text1 += f"  \u2014 {quote1.author}"
            text2 = f'\u201c{quote2.text}\u201d'
            if quote2.author:
                text2 += f"  \u2014 {quote2.author}"
            _tb1(slide,
                 Inches(3), SLIDE_H - Inches(2.2), SLIDE_W - Inches(6), Inches(0.7),
                 text1, FB, FS_QUOTE, WHITE, italic=True, align=PP_ALIGN.CENTER)
            _tb1(slide,
                 Inches(3), SLIDE_H - Inches(1.4), SLIDE_W - Inches(6), Inches(0.7),
                 text2, FB, FS_QUOTE, WHITE, italic=True, align=PP_ALIGN.CENTER)
        else:
            quote = quote_list[0]
            text = f'\u201c{quote.text}\u201d'
            if quote.author:
                text += f"  \u2014 {quote.author}"
            _tb1(slide,
                 Inches(3), SLIDE_H - Inches(1.6), SLIDE_W - Inches(6), Inches(1.4),
                 text, FB, FS_QUOTE, WHITE, italic=True, align=PP_ALIGN.CENTER)

_STATIC = {
    "truth": {
        "sidebar": "The Truth",
        "title": "The Truth About Your Pattern",
        "body": (
            "This isn't about fixing one of you.\n\n"
            "It's about changing the pattern\n"
            "that keeps pulling you apart."
        ),
        "say": (
            'Say: "Before we talk about logistics or sessions, I want to ground us '
            'in what this retreat is actually designed to change."'
        ),
        "quote": (
            "\u201cLove is not a feeling. Love is a practice of presence "
            "when the nervous system wants to flee.\u201d"
        ),
    },
    "getting_in_way": {
        "sidebar": "What\u2019s in the Way",
        "title": "What\u2019s Actually Getting in the Way",
        "body": (
            "The Real Issue:\n"
            "\u25c6  You\u2019re reacting from old nervous-system patterns, not the present moment\n"
            "\u25c6  Conversations escalate before either of you feels fully heard\n"
            "\u25c6  You\u2019ve never been given tools to interrupt the cycle together\n\n"
            "This isn\u2019t a communication problem.\n"
            "It\u2019s a regulation + pattern problem."
        ),
        "say": (
            "Say: \u201cMost couples try to talk their way out of a nervous-system issue. "
            "That\u2019s why insight alone hasn\u2019t been enough.\u201d"
        ),
        "quote": (
            "\u201cWhat triggers us in our partner is rarely about them \u2014 "
            "it\u2019s our nervous system recognizing an old threat that no longer exists.\u201d"
        ),
    },
    "how_retreat": {
        "sidebar": "How This Retreat",
        "title": "How This Retreat Creates Change",
        "body": (
            "Regulation & Safety\n"
            "Calm the nervous system so conversations don\u2019t spiral automatically\n\n"
            "Core Pattern Repair\n"
            "Heal the root dynamic driving repeated conflict\n\n"
            "Integration & Forward Path\n"
            "Build shared tools that actually work at home"
        ),
        "say": (
            'Say: "Every session fits into one of these three categories. '
            'Nothing here is random or extra."'
        ),
        "quote": (
            "\u201cHealing happens when we stop trying to think our way to safety "
            "and start creating it in our bodies.\u201d"
        ),
    },
    "journey": {
        "sidebar": "Your Journey",
        "title": "Your Transformation Journey",
        "body": (
            "Let me walk you through exactly what happens\n"
            "and why each piece matters"
        ),
        "quote": (
            "\u201cThis isn\u2019t about becoming different people. It\u2019s about interrupting "
            "the pattern long enough to remember who you really are together.\u201d"
        ),
    },
    "next_step": {
        "sidebar": "Your Next Step",
        "title": "Your Next Step",
        "body": (
            "If this feels aligned:\n"
            "\u25c6  We secure your dates\n"
            "\u25c6  Lock in your practitioner team for 12 months\n"
            "\u25c6  Begin pre-retreat intention work\n\n"
            "Place deposit today to hold availability"
        ),
        "say": (
            'Say: "What questions do you need answered to feel clear about moving forward?"'
        ),
        "quote": (
            "\u201cThe moment of decision is the beginning of transformation. "
            "Everything else is just preparation or delay.\u201d"
        ),
    },
}

def _s_static_rich(slide, key: str, quotes: Union[Optional[Quote], List[Quote], None] = None) -> None:
    data = _STATIC[key]
    title = data["title"]
    short = data.get("sidebar", title.split("—")[0].strip())

    _tb1(slide, SL_X, SL_Y, SL_W, SL_H, short, FT, FS_SIDEBAR, WHITE)

    _tb1(slide, CX, CY, CW, Inches(1.2), title, FT, FS_SLIDE_TITLE, ORANGE)
    
    quote_text = data.get("quote", "")
    if quotes:
        quote_list = []
        if isinstance(quotes, list):
            quote_list = quotes[:1]
        else:
            quote_list = [quotes]
        if quote_list and len(quote_list) >= 1:
            quote = quote_list[0]
            quote_text = f'\u201c{quote.text}\u201d'
            if quote.author:
                quote_text += f"  \u2014 {quote.author}"

    if key == "next_step":
        _tb1(slide, CX, CY + Inches(1.4), CW, Inches(0.8),
             quote_text, FB, FS_QUOTE, SAGE, italic=True, align=PP_ALIGN.CENTER)
        
        body_y  = CY + Inches(2.3)
        body_h  = PHOTO_Y - body_y - Inches(0.3)
        
        tf = _add_tb(slide, CX, body_y, CW, body_h)
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        first = True
        for line in data["body"].split("\n"):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            if not line.strip():
                tf.add_paragraph()
                continue
            _run(p, line, FB, FS_BODY, DARK)
        
        if "say" in data:
            say_text = data["say"]
    else:
        if key == "journey":
            quote_y = CY + Inches(1.4)
            if quote_text:
                _tb1(slide, CX, quote_y, CW, Inches(0.7),
                     quote_text, FB, FS_QUOTE, SAGE, italic=True, align=PP_ALIGN.CENTER)
                body_y = quote_y + Inches(0.9)
            else:
                body_y = quote_y
            
            body_text = data.get("body", "")
            tf = _add_tb(slide, CX, body_y, CW, Inches(1.0))
            lines = [line.strip() for line in body_text.split("\n") if line.strip()]
            for i, line in enumerate(lines):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = PP_ALIGN.CENTER
                r = p.add_run()
                r.text = line
                r.font.name = FB
                r.font.size = Pt(FS_BODY)
                r.font.color.rgb = DARK
        else:
            if key == "truth":
                quote_y = CY + Inches(1.4)
                if quote_text:
                    _tb1(slide, CX, quote_y, CW, Inches(0.7),
                         quote_text, FB, FS_QUOTE, SAGE, italic=True, align=PP_ALIGN.CENTER)
                    body_y = quote_y + Inches(0.9)
                else:
                    body_y = quote_y
                
                say_y   = PHOTO_Y - Inches(1.6)
                body_h  = say_y - body_y - Inches(1.4)

                tf = _add_tb(slide, CX, body_y, CW, body_h)
                tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
                body_lines = [l.strip() for l in data["body"].split("\n") if l.strip()]
                first = True
                i = 0
                while i < len(body_lines):
                    p = tf.paragraphs[0] if first else tf.add_paragraph()
                    first = False
                    line = body_lines[i]
                    if "changing the pattern" in line and i + 1 < len(body_lines) and "that keeps pulling you apart" in body_lines[i + 1]:
                        combined_line = f"{line} {body_lines[i + 1]}"
                        _run(p, combined_line, FB, FS_BODY, DARK)
                        i += 2
                    else:
                        _run(p, line, FB, FS_BODY, DARK)
                        i += 1

                if "say" in data:
                    _tb1(slide, CX, say_y, CW, Inches(1.2),
                         data["say"], FB, FS_SAY, SAGE, italic=True)
            else:
                if key == "getting_in_way":
                    quote_y = CY + Inches(1.4)
                    if quote_text:
                        _tb1(slide, CX, quote_y, CW, Inches(0.7),
                             quote_text, FB, FS_QUOTE, SAGE, italic=True, align=PP_ALIGN.CENTER)
                        body_y = quote_y + Inches(0.9)
                    else:
                        body_y = quote_y
                    
                    body_h = PHOTO_Y - body_y - Inches(0.3)

                    tf = _add_tb(slide, CX, body_y, CW, body_h)
                    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
                    body_lines = [l.strip() for l in data["body"].split("\n") if l.strip()]
                    first = True
                    i = 0
                    while i < len(body_lines):
                        p = tf.paragraphs[0] if first else tf.add_paragraph()
                        first = False
                        line = body_lines[i]
                        if "changing the pattern" in line and i + 1 < len(body_lines) and "that keeps pulling you apart" in body_lines[i + 1]:
                            combined_line = f"{line} {body_lines[i + 1]}"
                            _run(p, combined_line, FB, FS_BODY, DARK)
                            i += 2
                        else:
                            _run(p, line, FB, FS_BODY, DARK)
                            i += 1
                else:
                    if key == "how_retreat":
                        quote_y = CY + Inches(1.4)
                        if quote_text:
                            _tb1(slide, CX, quote_y, CW, Inches(0.7),
                                 quote_text, FB, FS_QUOTE, SAGE, italic=True, align=PP_ALIGN.CENTER)
                            body_y = quote_y + Inches(0.9)
                        else:
                            body_y = quote_y
                        
                        body_h = PHOTO_Y - body_y - Inches(0.3)

                        tf = _add_tb(slide, CX, body_y, CW, body_h)
                        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
                        heading_keys = {"Regulation & Safety", "Core Pattern Repair",
                                        "Integration & Forward Path", "The Real Issue:"}
                        first = True
                        for line in data["body"].split("\n"):
                            p = tf.paragraphs[0] if first else tf.add_paragraph()
                            first = False
                            if not line.strip():
                                tf.add_paragraph()
                                continue
                            is_heading = any(line.startswith(h) for h in heading_keys)
                            _run(p, line, FT if is_heading else FB,
                                 FS_BODY, ORANGE if is_heading else DARK,
                                 bold=is_heading)
                    else:
                        quote_y = CY + Inches(1.4)
                        if quote_text:
                            _tb1(slide, CX, quote_y, CW, Inches(0.7),
                                 quote_text, FB, FS_QUOTE, SAGE, italic=True, align=PP_ALIGN.CENTER)
                            body_y = quote_y + Inches(0.9)
                        else:
                            body_y = quote_y
                        
                        say_y   = PHOTO_Y - Inches(1.6)
                        body_h  = say_y - body_y - Inches(1.4)

                        tf = _add_tb(slide, CX, body_y, CW, body_h)
                        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
                        heading_keys = {"Regulation & Safety", "Core Pattern Repair",
                                        "Integration & Forward Path", "The Real Issue:"}
                        first = True
                        for line in data["body"].split("\n"):
                            p = tf.paragraphs[0] if first else tf.add_paragraph()
                            first = False
                            if not line.strip():
                                tf.add_paragraph()
                                continue
                            is_heading = any(line.startswith(h) for h in heading_keys)
                            _run(p, line, FT if is_heading else FB,
                                 FS_BODY, ORANGE if is_heading else DARK,
                                 bold=is_heading)

                        if "say" in data:
                            _tb1(slide, CX, say_y, CW, Inches(1.2),
                                 data["say"], FB, FS_SAY, SAGE, italic=True)

def _s_static(slide, title: str) -> None:
    _tb1(slide, CX, CY, CW, Inches(1.5), title, FT, FS_SLIDE_TITLE, ORANGE)
    short = title.split("—")[0].strip() if "—" in title else title
    _tb1(slide, SL_X, SL_Y, SL_W, SL_H, short, FT, FS_SIDEBAR, WHITE)

_SAY_WHAT_YOU_WANT = (
    'Say: \u201cIf this were happening consistently, '
    'how different would your relationship feel six months from now?\u201d'
)

def _s2_what_you_want(slide, bullets: List[str], quotes: Union[Optional[Quote], List[Quote]]) -> None:
    _tb1(slide, CX, CY, CW, Inches(1.3),
         "What You Told Me You Want Instead", FT, FS_SLIDE_TITLE, ORANGE)
    _tb1(slide, SL_X, SL_Y, SL_W, Inches(2.5),
         "What You Want", FT, FS_SIDEBAR, WHITE)

    quote_list = []
    if quotes:
        if isinstance(quotes, list):
            quote_list = quotes[:2]
        else:
            quote_list = [quotes]
    
    quote_y = CY + Inches(1.5)
    if quote_list and len(quote_list) >= 1:
        quote1 = quote_list[0]
        text1 = f'\u201c{quote1.text}\u201d'
        if quote1.author:
            text1 += f"  \u2014 {quote1.author}"
        _tb1(slide, CX, quote_y, CW, Inches(0.7),
             text1, FB, FS_QUOTE, SAGE, italic=True, align=PP_ALIGN.CENTER)
        body_start_y = quote_y + Inches(0.9)
    else:
        body_start_y = quote_y

    body_h = PHOTO_Y - body_start_y - Inches(0.2)

    tf = _add_tb(slide, CX, body_start_y, CW, body_h)
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    first = True
    for b in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        bullet_text = f"• {b}" if not b.startswith("•") else b
        _run(p, bullet_text, FB, FS_BODY, DARK)


def _s_day(slide, title: str, day: DayData,
           quotes: Union[Optional[Quote], List[Quote]], what_creates: List[str]) -> None:
    _tb1(slide, CX, CY, CW, Inches(1.2), title, FT, FS_SLIDE_TITLE, ORANGE)

    short = title.split("—")[0].strip() if "—" in title else title
    _tb1(slide, SL_X, SL_Y, SL_W, SL_H, short, FT, FS_SIDEBAR, WHITE)

    quote_list = []
    if quotes:
        if isinstance(quotes, list):
            quote_list = quotes[:2]
        else:
            quote_list = [quotes]
    
    quote_y = CY + Inches(1.4)
    if quote_list and len(quote_list) >= 1:
        quote1 = quote_list[0]
        text1 = f'\u201c{quote1.text}\u201d'
        if quote1.author:
            text1 += f"  \u2014 {quote1.author}"
        _tb1(slide, CX, quote_y, CW, Inches(0.7),
             text1, FB, FS_QUOTE, SAGE, italic=True, align=PP_ALIGN.CENTER)
        col_y = quote_y + Inches(0.9)
    else:
        col_y = quote_y

    col_h = PHOTO_Y - col_y - Inches(0.3)

    tf = _add_tb(slide, CX, col_y, COL_W, col_h)

    p = tf.paragraphs[0]
    _run(p, "Sessions", FT, FS_HEAD, ORANGE, bold=True)

    for s in day.sessions:
        _run(tf.add_paragraph(), f"• {s.title}", FB, FS_BODY, DARK)

    tf.add_paragraph()

    _run(tf.add_paragraph(), "Modalities", FT, FS_SMALL, SAGE, bold=True)
    if day.combined_modalities:
        _run(tf.add_paragraph(), day.combined_modalities, FB, FS_SMALL, SAGE,
             italic=True)

    tf2 = _add_tb(slide, COL2_X, col_y, COL_W, col_h)

    p2 = tf2.paragraphs[0]
    _run(p2, "What This Creates", FT, FS_HEAD, ORANGE, bold=True)

    for bullet in what_creates:
        _run(tf2.add_paragraph(), bullet, FB, FS_BODY, DARK)

def _s12_investment(slide, amount: str, details: List[str],
                    quotes: Union[Optional[Quote], List[Quote]]) -> None:
    _tb1(slide, CX, CY, CW, Inches(1.2),
         "Your Investment in Transformation", FT, FS_SLIDE_TITLE, ORANGE)
    _tb1(slide, SL_X, SL_Y, SL_W, Inches(2.0),
         "Investment", FT, FS_SIDEBAR, WHITE)

    quote_list = []
    if quotes:
        if isinstance(quotes, list):
            quote_list = quotes[:2]
        else:
            quote_list = [quotes]
    
    if quote_list and len(quote_list) >= 1:
        quote = quote_list[0]
        text = f'\u201c{quote.text}\u201d'
        if quote.author:
            text += f"  \u2014 {quote.author}"
        _tb1(slide, CX, CY + Inches(1.4), CW, Inches(0.7),
             text, FB, FS_QUOTE, SAGE, italic=True, align=PP_ALIGN.CENTER)
        amount_y = CY + Inches(2.2)
    else:
        amount_y = CY + Inches(1.4)

    amount_text = f"Total Investment: {amount}" if not amount.startswith("Total Investment") else amount
    _tb1(slide, CX, amount_y, CW, Inches(1.8),
         amount_text, FT, 60, DARK, bold=True)

    options_y = amount_y + Inches(2.0)
    if details:
        _tb1(slide, CX, options_y, CW, Inches(0.6),
             "Payment Options:", FT, FS_HEAD, SAGE, bold=True)
        
        tf = _add_tb(slide, CX, options_y + Inches(0.7), CW, Inches(1.2))
        first = True
        for line in details:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            if line.strip():
                bullet_text = line.strip()
                if not bullet_text.startswith("•") and not bullet_text.startswith("-"):
                    bullet_text = f"• {bullet_text}"
                _run(p, bullet_text, FB, FS_BODY, DARK)
        
        retreat_desc_y = options_y + Inches(2.0)
        has_description = any("private" in d.lower() or "customized" in d.lower() or "capacity" in d.lower() for d in details)
        if not has_description:
            _tb1(slide, CX, retreat_desc_y, CW, Inches(0.5),
                 "This retreat is private, customized, and capacity-limited", FB, FS_BODY, DARK, bold=True)
    else:
        retreat_desc_y = options_y
        _tb1(slide, CX, retreat_desc_y, CW, Inches(0.5),
             "This retreat is private, customized, and capacity-limited", FB, FS_BODY, DARK, bold=True)


class DeckGenerator:

    def generate(
        self,
        proposal: ProposalData,
        quotes: dict,
        pattern_bullets: List[str],
        what_creates: dict,
        output_path: Optional[str] = None,
    ) -> bytes:
        if not os.path.exists(TEMPLATE):
            raise FileNotFoundError(
                f"Template not found: {TEMPLATE}\n"
                "Place 'SSA PP template.pptx' in ppt_gen/"
            )

        prs = Presentation(TEMPLATE)

        for _ in range(12):
            _dup_slide(prs, 1)

        assert len(prs.slides) == 14, (
            f"Expected 14 slides after duplication, got {len(prs.slides)}"
        )

        slides = prs.slides

        _s0_title(slides[0], proposal.client_names, quotes.get("title"), proposal.retreat_title)
        logger.info("Slide 1: Title — %s", proposal.client_names)

        _s_static_rich(slides[1], "truth", quotes.get("truth"))
        logger.info("Slide 2: Truth (static)")

        _s2_what_you_want(slides[2], pattern_bullets, quotes.get("pattern"))
        logger.info("Slide 3: What You Want (%d bullets)", len(pattern_bullets))

        _s_static_rich(slides[3], "getting_in_way")
        _s_static_rich(slides[4], "how_retreat")
        _s_static_rich(slides[5], "journey")
        logger.info("Slides 4–6: static")

        if proposal.pre_retreat:
            _s_day(slides[6], "At-Home Pre-Retreat",
                   proposal.pre_retreat, quotes.get("pre_retreat"),
                   what_creates.get("pre_retreat", []))
            logger.info("Slide 7: Pre-Retreat")
        else:
            _s_static(slides[6], "At-Home Pre-Retreat")
            logger.warning("No pre-retreat data — slide 7 static")

        for i, day in enumerate(proposal.days[:4]):
            theme = day.theme or f"Day {day.day_number}"
            _s_day(
                slides[7 + i],
                f"Day {day.day_number} — {theme}",
                day,
                quotes.get(f"day_{day.day_number}"),
                what_creates.get(f"day_{day.day_number}", []),
            )
            logger.info("Slide %d: Day %d — %s", 8 + i, day.day_number, theme)

        if len(proposal.days) < 4:
            for i in range(len(proposal.days), 4):
                _s_static(slides[7 + i], f"Day {i + 1}")
            logger.warning(
                "Proposal has %d days; remaining day slides left static",
                len(proposal.days),
            )

        if proposal.post_retreat:
            _s_day(slides[11], "At-Home Integration & Phase II Support",
                   proposal.post_retreat, quotes.get("post_retreat"),
                   what_creates.get("post_retreat", []))
            logger.info("Slide 12: Post-Retreat")
        else:
            _s_static(slides[11], "At-Home Integration & Phase II Support")
            logger.warning("No post-retreat data — slide 12 static")

        _s12_investment(
            slides[12],
            amount=proposal.investment_amount or "$0",
            details=proposal.investment_details,
            quotes=quotes.get("investment"),
        )
        logger.info("Slide 13: Investment — %s", proposal.investment_amount)

        _s_static_rich(slides[13], "next_step")
        logger.info("Slide 14: Next Step (static)")

        buf = BytesIO()
        prs.save(buf)
        pptx_bytes = buf.getvalue()

        if output_path:
            with open(output_path, "wb") as f:
                f.write(pptx_bytes)
            logger.info("Saved: %s (%d bytes)", output_path, len(pptx_bytes))

        return pptx_bytes


